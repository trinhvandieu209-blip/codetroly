"""Xuất đề ra file Word để in cho học sinh làm trên giấy (kèm bản đáp án cho giáo viên)."""
import io
import random
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from services import images

SCHOOL_NAME = "Trường Liên cấp TH-THCS Nguyễn Quốc Phú"
IMAGE_WIDTH_CM = 9
MAX_ANSWER_LINES = 1


def _setup_document() -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.5)
        section.left_margin = section.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    return doc


def _header(doc: Document, quiz: dict[str, Any], with_answers: bool) -> None:
    school = doc.add_paragraph()
    school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    school.add_run(SCHOOL_NAME.upper()).bold = True
    subject = doc.add_paragraph(f"Môn Tin học lớp {quiz['grade']}")
    subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(quiz["title"] + (" – BẢN ĐÁP ÁN" if with_answers else ""))
    run.bold = True
    run.font.size = Pt(15)
    if not with_answers:
        doc.add_paragraph("Họ và tên: ………………………………………………   Lớp: …………   Điểm: ………")
    doc.add_paragraph(f"Đề gồm {len(quiz['questions'])} câu. Em đọc kỹ rồi khoanh tròn hoặc viết câu trả lời.").italic = True


def _shuffled_right(question: dict[str, Any]) -> list[tuple[str, str]]:
    """Vế phải xáo theo thứ tự cố định theo id câu, để bản đề và bản đáp án khớp nhau."""
    pairs = question["options"]
    order = list(range(len(pairs)))
    random.Random(question["id"]).shuffle(order)
    if order == sorted(order) and len(order) > 1:
        order = order[1:] + order[:1]
    return [(chr(ord("a") + k), pairs[idx][1]) for k, idx in enumerate(order)]


def _add_match(doc: Document, question: dict[str, Any]) -> None:
    pairs = question["options"]
    right = _shuffled_right(question)
    table = doc.add_table(rows=len(pairs), cols=2)
    table.style = "Table Grid"
    for i, pair in enumerate(pairs):
        table.cell(i, 0).text = f"{i + 1}. {pair[0]}"
        table.cell(i, 1).text = f"{right[i][0]}) {right[i][1]}"
    doc.add_paragraph("Trả lời: " + "   ".join(f"{i + 1} – ……" for i in range(len(pairs))))


def _add_image(doc: Document, question: dict[str, Any]) -> None:
    path = images.image_file(question.get("image_path", "")) if question.get("image_path") else None
    if path is None:
        return
    try:
        doc.add_picture(str(path), width=Cm(IMAGE_WIDTH_CM))
    except (OSError, ValueError):
        return


def _add_question(doc: Document, question: dict[str, Any], index: int) -> None:
    para = doc.add_paragraph()
    para.add_run(f"Câu {index}. ").bold = True
    para.add_run(question["prompt"])
    _add_image(doc, question)
    if question["qtype"] == "mcq":
        for opt in question["options"]:
            doc.add_paragraph(opt).paragraph_format.left_indent = Cm(1)
    elif question["qtype"] == "match":
        _add_match(doc, question)
    else:
        if question["options"]:
            doc.add_paragraph("Gợi ý từ: " + ", ".join(question["options"])).paragraph_format.left_indent = Cm(1)
        doc.add_paragraph("Trả lời: " + "…" * 60).paragraph_format.left_indent = Cm(1)


def _answer_text(question: dict[str, Any]) -> str:
    if question["qtype"] == "match":
        right = {text: letter for letter, text in _shuffled_right(question)}
        return "; ".join(f"{i + 1} – {right[pair[1]]}" for i, pair in enumerate(question["options"]))
    return question["answer"]


def _answer_key(doc: Document, quiz: dict[str, Any]) -> None:
    doc.add_page_break()
    head = doc.add_paragraph()
    head.add_run("ĐÁP ÁN VÀ GIẢI THÍCH").bold = True
    for i, q in enumerate(quiz["questions"], start=1):
        para = doc.add_paragraph()
        para.add_run(f"Câu {i}: ").bold = True
        para.add_run(_answer_text(q))
        if q.get("explanation"):
            para.add_run(f" — {q['explanation']}").italic = True


def build_quiz_docx(quiz: dict[str, Any], with_answers: bool = False) -> bytes:
    """Tạo file .docx của đề; with_answers=True thêm trang đáp án cho giáo viên."""
    doc = _setup_document()
    _header(doc, quiz, with_answers)
    for i, q in enumerate(quiz["questions"], start=1):
        _add_question(doc, q, i)
    if with_answers:
        _answer_key(doc, quiz)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
